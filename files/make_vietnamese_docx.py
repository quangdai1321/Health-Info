# -*- coding: utf-8 -*-
"""
Tạo bản Word (.docx) tiếng Việt đầy đủ của bài báo, dựng từ nội dung đã có
trong ban-dich-tieng-viet.md, trình bày theo cấu trúc bài báo khoa học
(tiêu đề, tóm tắt, các mục, Hình 1, Bảng 1, tài liệu tham khảo).

Đây là bản ĐỌC cho tác giả, không phải bản nộp hội nghị (bản nộp là
manuscript.tex bằng tiếng Anh), nên KHÔNG ẩn danh tác giả.
"""
import re
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

# --- Quy cach trinh bay (chuan luan van / bao cao khoa hoc VN) ---
FONT_NAME = 'Times New Roman'
FONT_SIZE_BODY = Pt(13)
FONT_SIZE_H1 = Pt(14)
FONT_SIZE_H2 = Pt(13)
LINE_SPACING = 1.5
PAGE_W, PAGE_H = Cm(21.0), Cm(29.7)          # A4
MARGIN_TOP, MARGIN_BOTTOM = Cm(2.0), Cm(2.0)
MARGIN_LEFT, MARGIN_RIGHT = Cm(3.0), Cm(2.0)  # trai 3cm chua cho dong gay
FIRST_LINE_INDENT = Cm(1.0)

SRC = "ban-dich-tieng-viet.md"
OUT = "HEALTHINFO-IV_TruongQuangDai_LLM-dien-giai-xet-nghiem-mau_TiengViet.docx"
FIG = "architecture.png"

INLINE_RE = re.compile(r'(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`]+?`)')


def _force_font(run, size=None):
    """Ep font ve Times New Roman cho ca phan chu Latin lan Unicode."""
    run.font.name = FONT_NAME
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rf)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(attr), FONT_NAME)
    if size is not None:
        run.font.size = size


def add_inline(paragraph, text, base_italic=False, size=None):
    """Thêm text vào paragraph, xử lý **đậm**, *nghiêng*, `code`."""
    size = size or FONT_SIZE_BODY
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            r = paragraph.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith('`') and tok.endswith('`'):
            # Ma/ky hieu ky thuat: giu cung font Times, chi in nghieng
            r = paragraph.add_run(tok[1:-1])
            r.italic = True
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1])
            r.italic = True
        else:
            r = paragraph.add_run(tok)
        if base_italic:
            r.italic = True
        _force_font(r, size)


def set_cell_text(cell, text, bold=False, align_center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if align_center
                   else WD_ALIGN_PARAGRAPH.LEFT)
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    if bold:
        r = p.add_run(text)
        r.bold = True
        _force_font(r, Pt(12))
    else:
        add_inline(p, text, size=Pt(12))


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'


def body_format(p, justify=True, indent=True):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = LINE_SPACING
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        pf.first_line_indent = FIRST_LINE_INDENT
    return p


def main():
    with open(SRC, encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()

    # --- Khổ giấy A4 + lề chuẩn ---
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = PAGE_W
        section.page_height = PAGE_H
        section.top_margin = MARGIN_TOP
        section.bottom_margin = MARGIN_BOTTOM
        section.left_margin = MARGIN_LEFT
        section.right_margin = MARGIN_RIGHT

    # --- Thống nhất font cho mọi style dùng trong bài ---
    for style_name, size, bold in (
            ('Normal', FONT_SIZE_BODY, None),
            ('Heading 1', FONT_SIZE_H1, True),
            ('Heading 2', FONT_SIZE_H2, True),
            ('List Bullet', FONT_SIZE_BODY, None),
            ('List Number', FONT_SIZE_BODY, None)):
        try:
            st = doc.styles[style_name]
        except KeyError:
            continue
        st.font.name = FONT_NAME
        st.font.size = size
        if bold is not None:
            st.font.bold = bold
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = rpr.makeelement(qn('w:rFonts'), {})
            rpr.append(rf)
        for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(attr), FONT_NAME)

    i = 0
    n = len(lines)
    title_done = False
    figure_inserted = False

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Tiêu đề bài (H1, chỉ 1 lần)
        if stripped.startswith('# ') and not title_done:
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.line_spacing = 1.3
            r = p.add_run(title_text.upper())
            r.bold = True
            _force_font(r, Pt(16))

            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run('ThS. Võ Hoàng Khang')
            r2.bold = True
            _force_font(r2, Pt(13))

            p3 = doc.add_paragraph()
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p3.paragraph_format.space_after = Pt(18)
            r3 = p3.add_run('Trường Đại học Công nghệ TP.HCM (HUTECH)')
            r3.italic = True
            _force_font(r3, Pt(13))

            title_done = True
            i += 1
            continue

        # Bỏ qua khối ghi chú (blockquote) và dòng ngang '---'
        if stripped.startswith('>') or stripped == '---' or stripped == '':
            i += 1
            continue

        # Mục lớn "## N. Tên mục"
        if stripped.startswith('## '):
            text = stripped[3:].strip()
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.line_spacing = 1.3
            add_inline(h, text, size=FONT_SIZE_H1)
            for r in h.runs:
                r.bold = True
            i += 1
            continue

        # Mục con "### ..."
        if stripped.startswith('### '):
            text = stripped[4:].strip()
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.line_spacing = 1.3
            add_inline(h, text, size=FONT_SIZE_H2)
            for r in h.runs:
                r.bold = True
                r.italic = False
            i += 1
            continue

        # Bảng markdown
        if stripped.startswith('|'):
            table_lines = []
            while i < n and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            rows = [
                [c.strip() for c in l.strip('|').split('|')]
                for l in table_lines
            ]
            # dòng thứ hai là dòng phân cách ---|---
            if len(rows) >= 2 and all(set(c) <= set('-: ') for c in rows[1]):
                header, data_rows = rows[0], rows[2:]
            else:
                header, data_rows = rows[0], rows[1:]
            table = doc.add_table(rows=1 + len(data_rows), cols=len(header))
            style_table(table)
            for c, htext in enumerate(header):
                set_cell_text(table.rows[0].cells[c], htext, bold=True,
                              align_center=True)
            for r_idx, row in enumerate(data_rows, start=1):
                for c, ctext in enumerate(row):
                    if c < len(table.columns):
                        set_cell_text(table.rows[r_idx].cells[c], ctext,
                                      align_center=(c > 0))
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # Danh sách gạch đầu dòng "- "
        if re.match(r'^-\s+', stripped):
            item_lines = [stripped]
            i += 1
            while i < n and lines[i].strip() and not re.match(
                    r'^(#|##|###|\||>|-\s|\d+\.\s)', lines[i].strip()):
                item_lines.append(lines[i].strip())
                i += 1
            full = ' '.join(l[2:].strip() if l.startswith('- ') else l
                             for l in item_lines)
            p = doc.add_paragraph(style='List Bullet')
            add_inline(p, full)
            body_format(p, justify=True, indent=False)
            continue

        # Danh sách đánh số "1. "
        m = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if m:
            item_lines = [m.group(2)]
            i += 1
            while i < n and lines[i].strip() and not re.match(
                    r'^(#|##|###|\||>|-\s|\d+\.\s)', lines[i].strip()):
                item_lines.append(lines[i].strip())
                i += 1
            full = ' '.join(item_lines)
            p = doc.add_paragraph(style='List Number')
            add_inline(p, full)
            body_format(p, justify=True, indent=False)
            continue

        # Đoạn văn thường: gộp các dòng liên tiếp thành 1 đoạn
        para_lines = [stripped]
        i += 1
        while i < n and lines[i].strip() and not re.match(
                r'^(#|##|###|\||>|-\s|\d+\.\s|---$)', lines[i].strip()):
            para_lines.append(lines[i].strip())
            i += 1
        full = ' '.join(para_lines)
        p = doc.add_paragraph()
        add_inline(p, full)
        # Chú thích bảng ("Bảng 1. ...") canh giữa, không thụt đầu dòng
        is_caption = full.startswith('**Bảng')
        body_format(p, justify=not is_caption, indent=not is_caption)
        if is_caption:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)

        # Chèn Hình 1 ngay sau câu dẫn "Hình 1 minh họa bốn tầng..."
        if not figure_inserted and full.startswith('Hình 1 minh họa'):
            doc.add_picture(FIG, width=Cm(14.0))
            last_p = doc.paragraphs[-1]
            last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            last_p.paragraph_format.space_before = Pt(6)
            last_p.paragraph_format.space_after = Pt(4)
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.line_spacing = 1.15
            cap.paragraph_format.space_after = Pt(12)
            r = cap.add_run(
                'Hình 1. Quy trình bốn tầng. Phán đoán lâm sàng (Tầng 2) '
                'hoàn tất trước khi mô hình ngôn ngữ (Tầng 3) được gọi; '
                'Tầng 4 kiểm chứng văn bản sinh ra so với nhãn của Tầng 2.'
            )
            r.italic = True
            _force_font(r, Pt(12))
            figure_inserted = True

    # --- Tài liệu tham khảo ---
    h = doc.add_heading(level=1)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    add_inline(h, 'Tài liệu tham khảo', size=FONT_SIZE_H1)
    for r in h.runs:
        r.bold = True
    # Danh muc IEEE, danh so theo THU TU TRICH DAN trong bai (khong xep A-Z).
    refs = [
        "T. D. Giardina, J. Baldwin, D. T. Nystrom, D. F. Sittig, and H. Singh, "
        "“Patient perceptions of receiving test results via online portals: "
        "A mixed-methods study,” Journal of the American Medical Informatics "
        "Association, vol. 25, no. 4, pp. 440–446, 2018.",

        "Z. He, B. Bhasuran, Q. Jin et al., “Quality of answers of generative "
        "large language models versus peer users for interpreting laboratory test "
        "results for lay patients: Evaluation study,” Journal of Medical "
        "Internet Research, vol. 26, p. e56655, 2024.",

        "A. Bozer and Y. Pekçevik, “Comparative evaluation of large "
        "language models in explaining radiology reports: Expert assessment of "
        "readability, understandability, and communication features,” "
        "Insights into Imaging, vol. 16, p. 232, 2025.",

        "R. Kianian, D. Sun, W. Rojas-Carabali, R. Agrawal, and E. Tsui, "
        "“Large language models may help patients understand peer-reviewed "
        "scientific articles about ophthalmology: Development and usability "
        "study,” Journal of Medical Internet Research, vol. 26, p. e59843, 2024.",

        "J. Cadamuro, F. Cabitza, Z. Debeljak et al., “Potentials and pitfalls "
        "of ChatGPT and natural-language artificial intelligence models for the "
        "understanding of laboratory medicine test results: An assessment by the "
        "European Federation of Clinical Chemistry and Laboratory Medicine (EFLM) "
        "Working Group on Artificial Intelligence,” Clinical Chemistry and "
        "Laboratory Medicine, vol. 61, no. 7, pp. 1158–1166, 2023.",

        "C. Munoz-Zuluaga, Z. Zhao, F. Wang, M. B. Greenblatt, and H. S. Yang, "
        "“Assessing the accuracy and clinical utility of ChatGPT in laboratory "
        "medicine,” Clinical Chemistry, vol. 69, no. 8, pp. 939–940, 2023.",

        "M. T. Tran, H. T. Lai, H. T. Hoang et al., “Health literacy among older "
        "adults: A descriptive cross-sectional analysis in Vietnam,” Journal of "
        "Health Literacy, vol. 10, no. 2, pp. 91–100, 2025.",

        "“Factors associated with health literacy among the elderly people in "
        "Vietnam,” BioMed Research International, vol. 2020, p. 3490635, 2020.",

        "T. T. H. Dinh and A. Bonner, “Psychometric properties of the health "
        "literacy questionnaire tested in Vietnamese adults with chronic "
        "diseases,” BMC Public Health, vol. 25, p. 44, 2025.",

        "B. N. Nguyen, T. T. Ngo et al., “Health and eHealth literacy in Vietnam: "
        "Evidence from a national survey,” Patient Education and Counseling, "
        "vol. 148, p. 109583, 2026.",

        "S. Geng, M. Josifoski, M. Peyrard, and R. West, “Grammar-constrained "
        "decoding for structured NLP tasks without finetuning,” in Proc. 2023 "
        "Conf. Empirical Methods in Natural Language Processing (EMNLP), 2023.",

        "B. T. Willard and R. Louf, “Efficient guided generation for large "
        "language models,” arXiv preprint arXiv:2307.09702, 2023.",

        "S. Choi, T. Fang, Z. Wang, and Y. Song, “KCTS: Knowledge-constrained "
        "tree search decoding with token-level hallucination detection,” in "
        "Proc. 2023 Conf. Empirical Methods in Natural Language Processing "
        "(EMNLP), 2023.",

        "P. Naliyatthaliyazchayil and T. Stenerson, “Harmonizing Logical "
        "Observation Identifiers Names and Codes (LOINC) codes and units in "
        "real-world oncology data: Method development and evaluation,” JMIR "
        "Medical Informatics, vol. 14, p. e81254, 2026.",

        "A. M. Zayed, I. Sarikakis, and N. Delvaux, “Automated standardization "
        "and harmonization of laboratory units in large-scale clinical data using "
        "open-source R functions,” International Journal of Medical "
        "Informatics, vol. 205, p. 106131, 2025.",

        "S. K. Parr, M. S. Shotwell, A. D. Jeffery, T. A. Lasko, and M. E. Matheny, "
        "“Automated mapping of laboratory tests to LOINC codes using noisy "
        "labels in a national electronic health record system database,” "
        "Journal of the American Medical Informatics Association, vol. 25, no. 10, "
        "pp. 1292–1300, 2018.",

        "J. de la Torre, “Scalable unit harmonization in medical informatics via "
        "Bayesian-optimized retrieval and transformer-based re-ranking,” "
        "International Journal of Medical Informatics, vol. 206, p. 106180, 2026.",
    ]
    for idx, ref in enumerate(refs, start=1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(1.0)
        pf.first_line_indent = Cm(-1.0)
        pf.line_spacing = 1.2
        pf.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, f"[{idx}] {ref}", size=Pt(12))

    doc.save(OUT)
    print("Da luu", OUT)


if __name__ == '__main__':
    main()
